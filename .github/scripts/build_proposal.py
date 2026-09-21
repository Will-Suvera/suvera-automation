#!/usr/bin/env python3
"""Build the post-call Recall proposal for one Partner Meeting Library page.

Runs after the meeting pipeline. For a Prospect meeting it:
  1. builds the proposal .docx from the template kept in Drive
     ("Suvera Proposals/_automation/Recall Proposal template.docx", a copy of a
     current single-practice proposal) - names, list size, pricing, ROI;
  2. uploads it as a Google Doc in "Suvera Proposals", editable by anyone at
     suvera.co.uk, and exports a PDF;
  3. puts a link-viewable copy of the PDF in "_automation" for the Dock builder
     to upload onto the workspace's Investment Proposal page;
  4. posts the Doc link in the meeting's Slack thread;
  5. writes Proposal Doc / Proposal PDF / Attendee Emails / Follow-up Email to
     the Notion page (the Dock routine and post-call-followup.yml read them).

Optional spec (--spec, written by the pipeline's Claude step) adds call detail:
  {"written_for": [{"name","role","email"}], "system": "EMIS", "local_scheme":
   "Enhanced Care Framework", "extra_bullets": [...], "short_name": "...",
   "list_size": 9300 (as stated on the call - beats the public registry),
   "practices": [{"name","ods","list_size"}] (several sites or a PCN: one
   pricing row per practice), "pcn_discount_agreed": true (the Suvera speaker
   agreed to honour the PCN rate), "agreed_terms": ["..."] (anything
   commercial agreed on the call beyond list price), "email_body": "...
   {DOCK_LINK} ...", "skip": false}

Pricing rules (Will, 17-18 Sep 2026): monthly fees, PCN discount rows in the
table, every figure to 2 decimal places (the rate shows as £0.68, not £0.675).
A Prospect call that cannot get a proposal is reported in the Slack thread.

Usage: build_proposal.py (--page <notion page id> | --rid <Meeting ID>) [--spec f.json] [--dry-run]
                         [--force] [--template local.docx] [--strict]
Env:   NOTION_TOKEN, SLACK_BOT_TOKEN, FATHOM_<ACCOUNT>_KEY; rclone "gdrive:" remote.
Never fails the job unless --strict.
"""
import copy
import json
import os
import re
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP

CHANNEL = "C0APW8DSA4R"
DS = "8115f0ee-00c8-488a-a05b-57726af0acf4"  # Partner Meeting Library
TEMPLATE_REMOTE = "gdrive:Suvera Proposals/_automation/Recall Proposal template.docx"
DOCS_FOLDER, PDF_SUBFOLDER = "Suvera Proposals", "_automation"
DOMAIN = "suvera.co.uk"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
GDOC_MIME = "application/vnd.google-apps.document"
FOLDER_MIME = "application/vnd.google-apps.folder"
WRITERS = {"will": ("Will Gao", "Founder & CCO: Suvera", "will@suvera.co.uk"),
           "ellena": ("Ellena Rollason", "Partnerships Manager: Suvera", "ellena.rollason@suvera.co.uk")}
# per-1,000 registered patients (same model as the recall-proposal skill)
ROI = [(270, 15, "{} admin hrs"), (540, 30, "{} admin hrs"), (2500, 100, "{} fewer appts*"),
       (390, 45, "{} fewer blood tests"), (1550, 62, "{} fewer appts"),
       (680, None, "{} QOF pts + {} extra reviews"), (600, 24, "{} appts recovered")]
SUFFIXES = ("group practice", "medical practice", "medical centre", "health centre",
            "family practice", "surgery", "practice")


def log(*a):
    print("[proposal]", *a, flush=True)


def http(method, url, body=None, headers=None, raw=False, retries=5):
    for attempt in range(retries):
        req = urllib.request.Request(url, data=body, method=method,
                                     headers={"User-Agent": "Mozilla/5.0 (suvera-automation)", **(headers or {})})
        try:
            with urllib.request.urlopen(req, timeout=90) as r:
                data = r.read()
            return data if raw else (json.loads(data) if data else {})
        except urllib.error.HTTPError as e:
            if e.code in (403, 429, 500, 502, 503) and attempt < retries - 1 and "googleapis" in url:
                time.sleep(8 * (attempt + 1))
                continue
            raise RuntimeError(f"{method} {url.split('?')[0]} -> {e.code}: {e.read().decode()[:300]}")


def notion(method, path, body=None):
    return http(method, "https://api.notion.com/v1" + path,
                None if body is None else json.dumps(body).encode(),
                {"Authorization": "Bearer " + os.environ["NOTION_TOKEN"],
                 "Notion-Version": "2025-09-03", "Content-Type": "application/json"})


def ptext(p):
    if not p:
        return ""
    t = p["type"]
    v = p.get(t)
    if t in ("title", "rich_text"):
        return "".join(x.get("plain_text", "") for x in v)
    if t == "select":
        return (v or {}).get("name", "")
    if t == "date":
        return (v or {}).get("start", "")
    return v or "" if isinstance(v, str) else ""


# ---------- Drive (via the rclone "gdrive:" remote's OAuth token) ----------
def drive_token():
    subprocess.run(["rclone", "about", "gdrive:"], capture_output=True)
    cfg = json.loads(subprocess.run(["rclone", "config", "dump"], capture_output=True, text=True).stdout)
    return json.loads(cfg["gdrive"]["token"])["access_token"]


def gapi(tok, method, url, body=None, ctype=None, raw=False):
    h = {"Authorization": "Bearer " + tok}
    if ctype:
        h["Content-Type"] = ctype
    return http(method, url, body, h, raw=raw)


def gquery(tok, q):
    return gapi(tok, "GET", "https://www.googleapis.com/drive/v3/files?" + urllib.parse.urlencode(
        {"q": q, "fields": "files(id,name,mimeType,webViewLink)"}))["files"]


def folder_id(tok, name, parent=None):
    q = f"name = '{name}' and mimeType = '{FOLDER_MIME}' and trashed = false"
    if parent:
        q += f" and '{parent}' in parents"
    f = gquery(tok, q)
    if f:
        return f[0]["id"]
    meta = {"name": name, "mimeType": FOLDER_MIME, **({"parents": [parent]} if parent else {})}
    return gapi(tok, "POST", "https://www.googleapis.com/drive/v3/files",
                json.dumps(meta).encode(), "application/json")["id"]


def upload(tok, meta, blob, mime, existing_id=None):
    b = uuid.uuid4().hex
    body = (f"--{b}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{json.dumps(meta)}\r\n"
            f"--{b}\r\nContent-Type: {mime}\r\n\r\n").encode() + blob + f"\r\n--{b}--".encode()
    ctype = f"multipart/related; boundary={b}"
    base = "https://www.googleapis.com/upload/drive/v3/files"
    if existing_id:
        return gapi(tok, "PATCH", f"{base}/{existing_id}?uploadType=multipart&fields=id,webViewLink", body, ctype)
    return gapi(tok, "POST", f"{base}?uploadType=multipart&fields=id,webViewLink", body, ctype)


def share(tok, fid, perm):
    try:
        gapi(tok, "POST", f"https://www.googleapis.com/drive/v3/files/{fid}/permissions?sendNotificationEmail=false",
             json.dumps(perm).encode(), "application/json")
    except RuntimeError as e:
        log("share warning:", e)


# ---------- inputs ----------
def lookup_list_size(ods, practice):
    for q in filter(None, [ods, practice]):
        try:
            d = http("GET", "https://apptconfig.suvera.com/lookup?" + urllib.parse.urlencode({"q": q}))
        except Exception:
            continue
        for m in d.get("matches", []):
            if (ods and m.get("ods") == ods) or (not ods and m.get("name", "").lower() == practice.lower()):
                return int(m.get("listSize") or 0)
    return 0


def fathom_invitees(account, rid, date_iso):
    key = os.environ.get(f"FATHOM_{account.upper()}_KEY", "")
    if not key or not rid.isdigit():
        return []
    try:
        start = datetime.fromisoformat(date_iso.replace("Z", "+00:00"))
    except ValueError:
        return []
    after = (start - timedelta(days=1)).strftime("%Y-%m-%dT%H:%M:%SZ")
    cursor = None
    for _ in range(10):
        q = {"created_after": after, **({"cursor": cursor} if cursor else {})}
        d = http("GET", "https://api.fathom.ai/external/v1/meetings?" + urllib.parse.urlencode(q),
                 headers={"X-Api-Key": key})
        for m in d.get("items", []):
            if str(m.get("recording_id")) == rid:
                return [i for i in (m.get("calendar_invitees") or []) if i.get("is_external")]
        cursor = d.get("next_cursor")
        if not cursor:
            break
    return []


def hubspot_title(email):
    """Job title from HubSpot for an attendee email, or ''."""
    tok = os.environ.get("HUBSPOT_TOKEN", "")
    if not (tok and email):
        return ""
    try:
        r = http("POST", "https://api-eu1.hubapi.com/crm/v3/objects/contacts/search",
                 json.dumps({"filterGroups": [{"filters": [{"propertyName": "email", "operator": "EQ", "value": email}]}],
                             "properties": ["jobtitle"], "limit": 1}).encode(),
                 {"Authorization": "Bearer " + tok, "Content-Type": "application/json"})
        return ((r.get("results") or [{}])[0].get("properties") or {}).get("jobtitle") or ""
    except Exception as e:
        log("hubspot title lookup failed:", e)
        return ""


def slack_note(ts, msg):
    if not (ts and os.environ.get("SLACK_BOT_TOKEN")):
        return
    try:
        http("POST", "https://slack.com/api/chat.postMessage",
             json.dumps({"channel": CHANNEL, "thread_ts": ts, "text": msg, "unfurl_links": False}).encode(),
             {"Authorization": "Bearer " + os.environ["SLACK_BOT_TOKEN"], "Content-Type": "application/json; charset=utf-8"})
    except Exception as e:
        log("slack warning:", e)


def person_name(inv):
    n, email = (inv.get("name") or "").strip(), inv.get("email") or ""
    if not n or "@" in n:
        parts = [re.sub(r"\d+", "", x) for x in re.split(r"[._-]", email.split("@")[0])]
        return " ".join(p.capitalize() for p in parts if p) or email
    m = re.match(r"^([^,(]+),\s*([^(]+?)\s*(\(.*\))?$", n)
    if m:
        return f"{m.group(2).strip().title()} {m.group(1).strip().title()}"
    return n


def short_of(practice):
    s = practice
    for suf in SUFFIXES:
        if s.lower().endswith(" " + suf):
            s = s[: -len(suf) - 1]
            break
    return s.strip() or practice


# ---------- document ----------
def r2(x):
    return Decimal(x).quantize(Decimal("0.01"), ROUND_HALF_UP)


def gbp(x):
    return f"£{r2(x):,}"


def whole(x):
    return int(Decimal(x).quantize(Decimal("1"), ROUND_HALF_UP))


RATES = [("2-year term (standard)", "£0.75", Decimal("0.75")),
         ("3-year term", "£0.68 (-10%)", Decimal("0.675")),
         ("With PCN discount - 2-year term", "£0.68 (-10%)", Decimal("0.675")),
         ("With PCN discount - 3-year term", "£0.60 (-20%)", Decimal("0.60"))]


def monthly(n, rate):
    return r2(Decimal(n) * rate / 12)


def build_docx(template, out, P, SHORT, N, pcn, system, local_scheme, extras, people, writer, setup_extra,
               practices=None, pcn_agreed=False, agreed_terms=()):
    """practices: [(name, list_size)] when the deal covers several sites; N is then their sum."""
    import docx
    from docx.text.paragraph import Paragraph
    d = docx.Document(template)
    k = Decimal(N) / 1000
    m2 = monthly(N, Decimal("0.675"))
    multi = practices and len(practices) > 1

    def set_par(p, text):
        runs = p.runs
        if not runs:
            p.add_run(text)
            return
        m = re.match(r"^(•\s*)", runs[0].text)
        if m and runs[0].text.strip() == "•" and len(runs) > 1:
            runs[1].text = text
            for r in runs[2:]:
                r.text = ""
        else:
            runs[0].text = (m.group(1) if m else "") + text
            for r in runs[1:]:
                r.text = ""

    def set_labelled(p, label, rest):
        runs = p.runs
        b = next(i for i, r in enumerate(runs) if r.bold and r.text.strip())
        runs[b].text = label
        after = runs[b + 1:]
        if after:
            after[0].text = " " + rest
            for r in after[1:]:
                r.text = ""
        else:
            p.add_run(" " + rest)

    def cell(c, text):
        ps = c.paragraphs
        set_par(ps[0], text)
        for extra in ps[1:]:
            extra._p.getparent().remove(extra._p)

    def find(start):
        return next(p for p in d.paragraphs if p.text.lstrip("• ").startswith(start))

    if multi:
        names = ", ".join(n for n, _ in practices[:-1]) + f" and {practices[-1][0]}"
        set_par(find("Suvera is delighted"),
                f"Suvera is delighted to share a proposal for {P} to use Recall to automate and streamline recalls across "
                f"{names}" + (" - with our PCN rate honoured for the network, as discussed." if pcn_agreed else
                              (f", with the PCN rate shown alongside for when {pcn} comes aboard." if pcn else ".")))
    else:
        set_par(find("Suvera is delighted"),
                f"Suvera is delighted to share a proposal for {P} to use Recall to automate and streamline recalls across the practice"
                + (f" - with the option to bring in the wider {pcn}." if pcn else "."))
    set_par(find("As a partner"), f"As a partner, {P} receives:")
    bullets, p = [], find("Full Recall platform access")
    while p is not None and p.text.strip().startswith("•"):
        bullets.append(p)
        nxt = p._p.getnext()
        p = Paragraph(nxt, p._parent) if nxt is not None and nxt.tag.endswith("}p") else None
    sized = (", ".join(f"{n} ({s:,})" for n, s in practices) if multi else f"{N:,}")
    new = [f"Full Recall platform access for all registered patients - {sized}",
           f"Daily clinical data extracts and refresh, fully {system} integrated",
           f"QOF, {local_scheme or 'LES'} and CQC performance tracking and recall list generation",
           "High-risk drug and DMARD monitoring built in (MHRA / BNF frequencies)",
           "Automated ICE/TQuest forms for every test due - plus those coming up, on one form",
           "One combined, personalised recall per patient - no more condition-by-condition invites",
           *[e for e in extras if e][:3],
           "Outbound patient SMS via Suvera, included in the platform fee",
           "Dedicated Support Team"]
    while len(bullets) < len(new):
        c = copy.deepcopy(bullets[-1]._p)
        bullets[-1]._p.addnext(c)
        bullets.append(Paragraph(c, bullets[-1]._parent))
    for b, t in zip(bullets, new):
        set_par(b, t)
    for b in bullets[len(new):]:
        b._p.getparent().remove(b._p)
    set_par(find("Platform fee is billed"),
            "Platform fee is billed per registered patient per year, unlimited SMS and email included, invoiced monthly. "
            "There are no set-up or onboarding fees. Monthly fees" + (" per practice" if multi else "") + ", ex-VAT:")
    set_par(find("*Patient count"),
            "*Patient count: indicative fees are based on "
            + (f"approximate registered list sizes of " + ", ".join(f"{s:,} ({n})" for n, s in practices) + f" - {N:,} in total."
               if multi else f"an approximate registered list size of {N:,}.")
            + " Actual fees are calculated against the registered patient count loaded into Recall at go-live, charged at the rate above.")
    who = pcn or "your PCN"
    set_par(find("PCN discount:"),
            ("PCN discount: as discussed, we are happy to honour our 10% PCN discount for " + who + " when the whole network comes aboard together"
             if pcn_agreed else
             "PCN discount: our 10% PCN discount applies when " + who + " comes aboard as a whole network")
            + (" - every practice, including " + SHORT + ", moves to the PCN rate." if not multi else ".")
            + " The two-year term is standard; a three-year commitment earns a further 10%, taking the rate to £0.60 per patient (20% off). "
            "Rates are shown to the nearest penny; discounts are applied to the £0.75 base rate before rounding.")
    if agreed_terms:
        anchor = find("PCN discount:")
        c = copy.deepcopy(anchor._p)
        anchor._p.addnext(c)
        agreed = Paragraph(c, anchor._parent)
        set_par(agreed, "Agreed on the call: " + " ".join(t.strip().rstrip(".") + "." for t in agreed_terms if t.strip()))
    set_labelled(find("Introduce a practice"), "Introduce a practice.",
                 f"For each practice you introduce that signs up, both {SHORT} and that practice receive one month off their invoice.")
    set_labelled(find("Bring in the PCN"), "Bring in the PCN.",
                 f"If {SHORT} brings in a PCN-wide contract for {pcn or 'your PCN'}, {SHORT} also receives a further one month "
                 f"off its invoice - worth {gbp(m2)} + VAT at the 2-year PCN rate.")
    set_par(find("Setup takes around"),
            "Setup takes around three weeks: clinical system and ICE access and integration, "
            + (setup_extra.rstrip(", ") + ", " if setup_extra else "")
            + f"then an in-person session to configure appointment slot types to your clinicians' competencies. "
            f"From there, Recall runs in the background at {SHORT}.")

    t_for, t_by, t_price, t_roi = d.tables[:4]
    rows = t_for.rows[1:]
    while len(t_for.rows) - 1 < len(people):
        t_for._tbl.append(copy.deepcopy(t_for.rows[-1]._tr))
    for r in t_for.rows[1 + len(people):]:
        t_for._tbl.remove(r._tr)
    for r, ppl in zip(t_for.rows[1:], people):
        for c, v in zip(r.cells, ppl):
            cell(c, v)
    for c, v in zip(t_by.rows[1].cells, writer):
        cell(c, v)
    if multi:
        # one row per practice + total; the four columns become the four rates
        price_rows = [(f"{n} ({s:,})", *[gbp(monthly(s, r)) for _, _, r in RATES]) for n, s in practices]
        price_rows.append((f"{P} total ({N:,})", *[gbp(sum(monthly(s, r) for _, s in practices)) for _, _, r in RATES]))
        # the template table has 4 columns: Practice | 2-year | 3-year | PCN 3-year
        # (the PCN 2-year rate equals the 3-year rate, so it needs no column of its own)
        cols = [0, 1, 2, 4]
        header = ("Practice (registered patients*)", "List price £0.75 (2-year term)",
                  "£0.68, -10% (3-year term)", "PCN rate £0.68, -10% (2-year term)", "PCN rate £0.60, -20% (3-year term)")
        price_rows = [tuple(row[i] for i in cols) for row in price_rows]
        header = tuple(header[i] for i in cols)
        for c, x in zip(t_price.rows[0].cells, header):
            cell(c, x)
    else:
        price_rows = [(f"{P} - {lab}" if not lab.startswith("With PCN") else lab.replace("With PCN discount", f"With PCN discount (if {who} joins)" if i == 2 else "With PCN discount"),
                       rate, gbp(monthly(N, r)), gbp(monthly(N, r) * Decimal("1.2")))
                      for i, (lab, rate, r) in enumerate(RATES)]
    while len(t_price.rows) - 1 < len(price_rows):
        t_price._tbl.append(copy.deepcopy(t_price.rows[1]._tr))
    for r in t_price.rows[1 + len(price_rows):]:
        t_price._tbl.remove(r._tr)
    for r, v in zip(t_price.rows[1:], price_rows):
        for c, x in zip(r.cells, v):
            cell(c, x)
    if multi:
        for c in t_price.rows[-1].cells:
            for run in c.paragraphs[0].runs:
                run.bold = True
    n = lambda v: f"{whole(Decimal(v) * k):,}"
    total = 0
    for i, (val, cap, fmt) in enumerate(ROI, start=1):
        v = whole(Decimal(val) * k)
        total += v
        text = fmt.format(n(2), n(10)) if cap is None else fmt.format(n(cap))
        cell(t_roi.rows[i].cells[1], text)
        cell(t_roi.rows[i].cells[2], f"£{v:,}")
    appts = sum(whole(Decimal(x) * k) for x in (100, 62, 24))
    hrs = sum(whole(Decimal(x) * k) for x in (15, 30))
    fee = whole(Decimal(N) * Decimal("0.75"))
    cell(t_roi.rows[8].cells[1], f"~{appts:,} appointments + {hrs:,} admin hrs")
    cell(t_roi.rows[8].cells[2], f"£{total:,}")
    cell(t_roi.rows[9].cells[2], f"(£{fee:,})")
    cell(t_roi.rows[10].cells[0], f"Indicative net value to {P} (Year 1)")
    cell(t_roi.rows[10].cells[2], f"£{total - fee:,} (~{(Decimal(total) / fee).quantize(Decimal('0.1'))}× cost)")
    # Google's PDF export ignores keep-with-next, so start Pricing on a fresh page to keep its table whole
    next(p for p in d.paragraphs if p.text.strip() == "Pricing").paragraph_format.page_break_before = True
    d.save(out)
    check = docx.Document(out)
    blob = "\n".join(p.text for p in check.paragraphs) + "\n".join(
        c.text for t in check.tables for r in t.rows for c in r.cells)
    for leftover in ("Waterfield", "Tunbridge", "Gonsalves", "Justice", "Minkah", "Stainfield", "Trehan"):
        if leftover in blob:
            raise RuntimeError(f"template text left in proposal: {leftover}")
    return blob


def main(argv):
    arg = lambda f, default=None: argv[argv.index(f) + 1] if f in argv else default
    dry, force = "--dry-run" in argv, "--force" in argv
    spec = {}
    if arg("--spec") and os.path.exists(arg("--spec")):
        try:
            spec = json.load(open(arg("--spec")))
        except ValueError as e:
            log("spec unreadable, using defaults:", e)
    if spec.get("skip"):
        log("spec says skip:", spec.get("reason", ""))
        return
    page_id = arg("--page") or spec.get("notion_page_id")
    if not page_id and arg("--rid"):
        res = notion("POST", f"/data_sources/{DS}/query",
                     {"filter": {"property": "Meeting ID", "rich_text": {"equals": arg("--rid")}}})["results"]
        page_id = res[0]["id"] if res else None
    if not page_id:
        log("no Notion page for this meeting - skipping")
        return
    page = notion("GET", f"/pages/{page_id}")
    pr = page["properties"]
    P = ptext(pr.get("Practice")) or spec.get("practice", "")
    if not P:
        # Titles look like "The Albion Surgery — Emily Randall (16 September 2026)"
        title = ptext(pr.get("Meeting"))
        P = re.split(r"\s+[\u2014\u2013-]\s+", title, maxsplit=1)[0].strip() if title else ""
        if P:
            log("Practice empty on the page; using the meeting title:", P)
    stage, account = ptext(pr.get("Stage")), (ptext(pr.get("Fathom Account")) or "will").lower()
    if not P:
        log("no Practice on the page - skipping")
        return
    if stage != "Prospect" and not force:
        log(f"stage is {stage!r}, not Prospect - no proposal")
        return
    if ptext(pr.get("Proposal Doc")) and not force:
        log("proposal already made:", ptext(pr.get("Proposal Doc")))
        return
    ods, pcn = ptext(pr.get("ODS Code")), ptext(pr.get("PCN"))
    ts = ptext(pr.get("Slack TS"))
    # Several sites or a PCN: one pricing row per practice. Registry sizes fill any the call did not state.
    practices = []
    for x in spec.get("practices") or []:
        name = (x.get("name") or "").strip()
        size = x.get("list_size") or lookup_list_size(x.get("ods") or "", name)
        if name and size:
            practices.append((name, int(size)))
    # A stated list size beats the public registry (Godolphin said 9,300; the registry had 8,931)
    N = spec.get("list_size") or (sum(s for _, s in practices) if len(practices) > 1 else 0) or lookup_list_size(ods, P)
    if not N:
        log("no list size (ODS lookup failed) - skipping")
        if stage == "Prospect":
            slack_note(ts, f":warning: No proposal was built for *{P}*: no list size found (ODS lookup failed and none was stated on the call). "
                           "Run the 'Build proposal (manual)' workflow with a list size, or add the ODS code to the Notion page.")
        return
    date_iso = ptext(pr.get("Date")) or datetime.utcnow().isoformat()
    when = datetime.fromisoformat(date_iso.replace("Z", "+00:00"))
    date_label = f"{when.day} {when.strftime('%B %Y')}"
    invitees = fathom_invitees(account, ptext(pr.get("Meeting ID")), date_iso)
    # Claude sometimes writes agreed_terms as one string rather than a list (21 Sep 2026, Family Medical Centre)
    agreed_terms = spec.get("agreed_terms") or []
    if isinstance(agreed_terms, str):
        agreed_terms = [t.strip() for t in re.split(r"(?<=[.;])\s+", agreed_terms) if t.strip()]
    agreed_terms = [str(t) for t in agreed_terms if str(t).strip()]

    def role_for(email, given):
        if given and not given.startswith("Practice team"):
            return given
        title = hubspot_title(email)
        return f"{title}: {P}" if title else (given or f"Practice team: {P}")
    if spec.get("written_for"):
        people = [(x.get("name", ""), role_for(x.get("email", ""), x.get("role") or ""), x.get("email", "")) for x in spec["written_for"]]
    elif invitees:
        people = [(person_name(i), role_for(i.get("email", ""), ""), i.get("email", "")) for i in invitees]
    else:
        people = [("Practice team", P, "")]
    people = [p for p in people if p[0] and (" " in p[0].strip() or p[2])]  # drop bare first names with no email
    emails = [e for _, _, e in people if e] or [i.get("email") for i in invitees if i.get("email")]
    writer = WRITERS.get(account, WRITERS["will"])
    SHORT = spec.get("short_name") or short_of(P)
    name = f"{P} - Suvera Recall Proposal - {date_label}" + (arg("--name-suffix") or "")
    out = f"/tmp/{name}.docx"
    template = arg("--template")
    if not template:
        template = "/tmp/proposal_template.docx"
        r = subprocess.run(["rclone", "copyto", TEMPLATE_REMOTE, template], capture_output=True, text=True)
        if r.returncode:
            raise RuntimeError("template download failed: " + r.stderr[-300:])
    text = build_docx(template, out, P, SHORT, int(N), pcn, spec.get("system") or "EMIS and SystmOne",
                      spec.get("local_scheme", ""), spec.get("extra_bullets", []), people, writer, spec.get("setup_extra", ""),
                      practices=practices, pcn_agreed=bool(spec.get("pcn_discount_agreed")),
                      agreed_terms=agreed_terms)
    log("built", out, f"({N:,} patients, {len(people)} recipients)")
    first = [p[0].split()[0] for p in people if p[0] and p[0] != "Practice team"]
    greet = ", ".join(first) if 0 < len(first) <= 4 else "all"
    email = spec.get("email_body") or (
        f"Hi {greet},\n\nThanks again for your time today - really appreciated it.\n\n"
        "I've put everything we covered in one place for you here - including the proposal, pricing and a quick demo video: {DOCK_LINK}\n\n"
        "Let me know if you have any questions at all - happy to walk through it on a call.\n\nVery best,\n"
        + writer[0].split()[0])
    if "{DOCK_LINK}" not in email:
        email += "\n\n{DOCK_LINK}"
    if dry:
        print(text)
        print("EMAILS", emails)
        print("EMAIL\n" + email)
        return
    tok = drive_token()
    docs = folder_id(tok, DOCS_FOLDER)
    pdfs = folder_id(tok, PDF_SUBFOLDER, docs)
    blob = open(out, "rb").read()
    esc = name.replace("'", "\\'")
    ex = gquery(tok, f"name = '{esc}' and '{docs}' in parents and mimeType = '{GDOC_MIME}' and trashed = false")
    doc = upload(tok, {} if ex else {"name": name, "mimeType": GDOC_MIME, "parents": [docs]}, blob, DOCX_MIME,
                 ex[0]["id"] if ex else None)
    share(tok, doc["id"], {"type": "domain", "domain": DOMAIN, "role": "writer"})
    pdf = gapi(tok, "GET", f"https://www.googleapis.com/drive/v3/files/{doc['id']}/export?mimeType=application/pdf", raw=True)
    exp = gquery(tok, f"name = '{esc}.pdf' and '{pdfs}' in parents and trashed = false")
    pf = upload(tok, {} if exp else {"name": name + ".pdf", "parents": [pdfs]}, pdf, "application/pdf",
                exp[0]["id"] if exp else None)
    share(tok, pf["id"], {"type": "anyone", "role": "reader"})
    pdf_url = f"https://drive.google.com/uc?export=download&id={pf['id']}"
    log("doc", doc["webViewLink"], "| pdf", pdf_url)
    if ts and os.environ.get("SLACK_BOT_TOKEN") and "--no-slack" not in argv:
        terms = ("\nAgreed on the call: " + " | ".join(agreed_terms)) if agreed_terms else ""
        msg = (f":page_facing_up: *Proposal - {P}* (Google Doc, anyone at Suvera can edit)\n<{doc['webViewLink']}|{name}>\n"
               f"{int(N):,} patients" + (f" across {len(practices)} practices" if len(practices) > 1 else "")
               + f": £0.75 (2-year) / £0.68 (3-year) / PCN £0.68 and £0.60 per patient, monthly fees.{terms}\n"
               "Check the figures against the call before the client sees it. The Dock checklist follows in this thread once the workspace is built.")
        r = http("POST", "https://slack.com/api/chat.postMessage",
                 json.dumps({"channel": CHANNEL, "thread_ts": ts, "text": msg, "unfurl_links": False}).encode(),
                 {"Authorization": "Bearer " + os.environ["SLACK_BOT_TOKEN"], "Content-Type": "application/json; charset=utf-8"})
        log("slack", r.get("ok"), r.get("error", ""))
    rt = lambda s: {"rich_text": [{"type": "text", "text": {"content": s[:1990]}}]}
    props = {"Proposal Doc": {"url": doc["webViewLink"]}, "Proposal PDF": {"url": pdf_url},
             "Attendee Emails": rt(", ".join(emails)), "Follow-up Email": rt(email)}
    if "--no-notion" in argv:
        log("--no-notion: not writing", list(props))
        return
    try:
        notion("PATCH", f"/pages/{page_id}", {"properties": props})
    except RuntimeError as e:
        log("notion warning:", e)
    log("done")


if __name__ == "__main__":
    try:
        main(sys.argv[1:])
    except Exception as e:  # never break the meeting pipeline
        log("FAILED:", e)
        if "--strict" in sys.argv:
            raise
